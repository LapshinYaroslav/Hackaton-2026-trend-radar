"""Задача П0.3: share_ru и share_patent по отдельности — без сети, из кэша П0.2.

Таблица 160 строк: n_ru и n_pat из кэша сбора, n_oa (OpenAlex) и n_research
(OpenAlex + arXiv) — из кэша счётчиков обучения ровно по шести годовым окнам 2020–2025,
без prev6. Доли считают model.features.share_ru и share_patent.

AUC одного признака — по фолдам той же кросс-валидации (StratifiedGroupKFold 5 × 10,
сиды SEED + повтор): пропуск в проверочной части заполняется медианой обучающей,
AUC считается по сырому значению признака. Ниже 0.5 — признак выше у мейнстрима.

Листы дописываются в report_tables.xlsx; прежние листы не трогаются.

Запуск: python -m scripts.ru_patents_p0
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import roc_auc_score
from sklearn.model_selection import StratifiedGroupKFold

from collector.constants import AGGREGATE_WINDOWS
from collector.models import SearchTerms
from model.candidate import cache_fetch
from model.config import N_REPEATS, N_SPLITS, ROSPATENT_DATASETS, SEED
from model.features import RESEARCH_SOURCES, share_patent
from model.train import training_table
from scripts.ru_patents_collect import TECHNOLOGIES, cache_path, request_key

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report_tables.xlsx"
FEATURES_P = {"share_ru": "n_ru", "share_patent": "n_pat"}
SIX = list(AGGREGATE_WINDOWS["all"])
CLASSES = ((1, "сигналы"), (0, "мейнстрим"))


def share_ru(n_ru: int, n_oa: int) -> float:
    """Доля работ OpenAlex с российским автором. Признак остановлен в П0 и в model/ не входит;
    формула живёт здесь, чтобы отчёт П0 воспроизводился. NaN, если n_oa = 0."""
    return n_ru / n_oa if n_oa else float("nan")


def counts(row: pd.Series) -> dict:
    """Четыре счётчика одной технологии: n_ru, n_pat из кэша П0.2, n_oa и n_research из обучения."""
    phrase = json.loads(row["terms"])[0]
    oa = json.loads(cache_path(request_key("openalex", phrase)).read_text(encoding="utf-8"))
    rp = json.loads(cache_path(request_key("rospatent", phrase)).read_text(encoding="utf-8"))
    counters = cache_fetch(SearchTerms(terms=[phrase], context_terms=[], query=""))
    six = counters.loc[counters["window"].isin(SIX)]
    return {"tech_id": row["tech_id"], "n_ru": int(oa["response"]["meta"]["count"]),
            "n_pat": int(rp["response"]["total"]),
            "n_oa": int(six.loc[six["source"] == "openalex", "n"].sum()),
            "n_research6": int(six.loc[six["source"].isin(RESEARCH_SOURCES), "n"].sum())}


def feature_table() -> pd.DataFrame:
    """Обучающая таблица s2a1-v1 плюс четыре счётчика и две доли."""
    technologies = pd.read_csv(TECHNOLOGIES, encoding="utf-8-sig")
    extra = pd.DataFrame([counts(row) for _, row in technologies.iterrows()])
    base = training_table().drop(columns=["n_pat", "share_patent"])  # П0 считает их сам, из кэша
    frame = base.reset_index(drop=True).merge(extra, on="tech_id", how="left")
    if (frame["n_research6"] != frame["n_research"]).any():
        raise ValueError("n_research из кэша по шести окнам разошёлся с обучающей таблицей")
    frame.loc[:, "share_ru"] = [share_ru(a, b) for a, b in zip(frame["n_ru"], frame["n_oa"])]
    frame.loc[:, "share_patent"] = [share_patent(a, b) for a, b
                                    in zip(frame["n_pat"], frame["n_research6"])]
    return frame


def subsets(frame: pd.DataFrame) -> dict[int, pd.DataFrame]:
    """160 строк и 149 строк с blind_choice."""
    blind = frame.loc[frame["blind_choice"].astype(bool)].reset_index(drop=True)
    return {len(frame): frame, len(blind): blind}


def windows() -> pd.DataFrame:
    """Границы окон числителя и знаменателя каждого признака."""
    return pd.DataFrame([
        {"признак": "share_ru", "числитель": "n_ru: OpenAlex, фраза, type:article, country_code:RU, "
         "from/to_publication_date 2020-09-01…2026-08-31, один вызов",
         "знаменатель": "n_oa: OpenAlex, фраза, type:article, сумма годовых окон 2020–2025 "
         "(2020-09-01…2026-08-31), кэш обучения"},
        {"признак": "share_patent", "числитель": "n_pat: Роспатент total, фраза в кавычках, "
         "date_published 20200901…20260831, датасеты: " + ", ".join(ROSPATENT_DATASETS),
         "знаменатель": "n_pat + n_research: OpenAlex + arXiv, сумма годовых окон 2020–2025 "
         "(2020-09-01…2026-08-31), без prev6, кэш обучения"}])


def describe(frames: dict[int, pd.DataFrame]) -> pd.DataFrame:
    """По признаку, классу и выборке: доля ненулевых счётчиков, квартили счётчика и доли."""
    rows = []
    for count, frame in frames.items():
        for feature, counter in FEATURES_P.items():
            for label, title in CLASSES:
                block = frame.loc[frame["label"] == label]
                q_n = block[counter].quantile([0.25, 0.5, 0.75]).to_numpy()
                q_s = block[feature].quantile([0.25, 0.5, 0.75]).to_numpy()
                rows.append({"строк": count, "признак": feature, "класс": title,
                             "n": len(block), "доля счётчик > 0":
                             round(float((block[counter] > 0).mean()), 3),
                             "счётчик": counter, "счётчик Q1": q_n[0],
                             "счётчик медиана": q_n[1], "счётчик Q3": q_n[2], "доля Q1": round(q_s[0], 4),
                             "доля медиана": round(q_s[1], 4), "доля Q3": round(q_s[2], 4),
                             "пропусков доли": int(block[feature].isna().sum())})
    return pd.DataFrame(rows)


def cv_auc(frame: pd.DataFrame, column: str) -> tuple[float, float]:
    """AUC сырого признака вне обучения: пропуск — медиана обучающей части фолда.

    Среднее по фолдам внутри повтора, затем среднее и std по повторам. Фолд, где в
    проверочной части один класс, пропускается (бывает в малых терцилях).
    """
    labels, values = frame["label"].to_numpy(), frame[column].to_numpy(dtype=float)
    per_repeat = []
    for repeat in range(N_REPEATS):
        splitter = StratifiedGroupKFold(n_splits=N_SPLITS, shuffle=True,
                                        random_state=SEED + repeat)
        folds = []
        for train, test in splitter.split(values, labels, frame["group"]):
            if len(set(labels[test])) < 2:
                continue
            filled = np.where(np.isnan(values[test]), np.nanmedian(values[train]), values[test])
            folds.append(roc_auc_score(labels[test], filled))
        per_repeat.append(float(np.mean(folds)))
    return float(np.mean(per_repeat)), float(np.std(per_repeat))


def auc_table(frames: dict[int, pd.DataFrame]) -> pd.DataFrame:
    """AUC одного признака на 160 и 149 строках."""
    rows = []
    for count, frame in frames.items():
        for feature in FEATURES_P:
            mean, std = cv_auc(frame, feature)
            rows.append({"строк": count, "признак": feature, "AUC": round(mean, 3),
                         "std по повторам": round(std, 3)})
    return pd.DataFrame(rows)


def spearman_table(frames: dict[int, pd.DataFrame]) -> pd.DataFrame:
    """Спирмен признака с anchor_volume и volume внутри каждого класса."""
    rows = []
    for count, frame in frames.items():
        for feature in FEATURES_P:
            for label, title in CLASSES:
                block = frame.loc[frame["label"] == label]
                row = {"строк": count, "признак": feature, "класс": title}
                for other in ("anchor_volume", "volume"):
                    pair = block[[feature, other]].astype(float).dropna()
                    row[f"~ {other}"] = round(float(pair[feature].corr(pair[other],
                                                                       method="spearman")), 3)
                    row[f"n пар ~ {other}"] = len(pair)
                rows.append(row)
    return pd.DataFrame(rows)


def tertile_auc(frames: dict[int, pd.DataFrame]) -> pd.DataFrame:
    """AUC признака внутри терцилей anchor_volume (терцили по рангу внутри выборки)."""
    rows = []
    for count, frame in frames.items():
        tertile = pd.qcut(frame["anchor_volume"].rank(method="first"), 3,
                          labels=["T1 низкий", "T2 средний", "T3 высокий"])
        for name in tertile.cat.categories:
            block = frame.loc[tertile == name].reset_index(drop=True)
            bounds = block["anchor_volume"]
            for feature in FEATURES_P:
                mean, std = cv_auc(block, feature)
                rows.append({"строк": count, "терциль": name, "anchor_volume от": bounds.min(),
                             "anchor_volume до": bounds.max(), "n": len(block),
                             "сигналов": int(block["label"].sum()), "признак": feature,
                             "AUC": round(mean, 3), "std": round(std, 3)})
    return pd.DataFrame(rows)


def cross_spearman(frame: pd.DataFrame) -> pd.DataFrame:
    """Спирмен share_ru ~ share_patent на 160 строках, по парам без пропусков."""
    pair = frame[["share_ru", "share_patent"]].dropna()
    return pd.DataFrame([{"строк": len(frame), "n пар": len(pair),
                          "rho": round(float(pair["share_ru"].corr(pair["share_patent"],
                                                                  method="spearman")), 3)}])


# Замер П0.1 от 24.09.2026 (пробные запросы до массового сбора, в кэш не входят).
PROBE_ROSPATENT = [
    ("large language model", 856, 275267, 966, 856),
    ("digital twin", 8119, 24572, 10105, 1000),
    ("neural network", 282616, 347927, 653702, 1000),
    ("edge model compression", 0, 119991, 0, 0),
]
PROBE_OPENALEX = [
    ("s1", "edge model compression", 14, 14, 0, 0),
    ("s8", "neuromorphic chip", 1638, 1638, 35, 35),
    ("n1", "bluetooth low energy", 9795, 9807, 75, 75),
    ("n13", "web application firewall", 1665, 1665, 21, 21),
]


def probe_sheets() -> dict[str, pd.DataFrame]:
    """Таблицы проверки источников П0.1."""
    rospatent = pd.DataFrame(PROBE_ROSPATENT, columns=[
        "фраза", "в кавычках, окно", "без кавычек, окно", "в кавычках, без окна",
        "available (в кавычках, окно)"])
    openalex = pd.DataFrame(PROBE_OPENALEX, columns=[
        "tech_id", "фраза", "кэш n_oa (6 окон)", "сейчас, всё окно", "n_ru group_by (1 кредит)",
        "n_ru полная выдача (10 кредитов)"])
    return {"П0.1 Роспатент": rospatent, "П0.1 OpenAlex": openalex}


def stop_rule(desc: pd.DataFrame, auc: pd.DataFrame, rho: pd.DataFrame) -> pd.DataFrame:
    """Правило остановки на 160 строках: по условию на строку, итог — хоть одно выполнено."""
    rows = []
    for feature in FEATURES_P:
        nonzero = desc.loc[(desc["строк"] == 160) & (desc["признак"] == feature)]
        worst_rho = rho.loc[(rho["строк"] == 160) & (rho["признак"] == feature), "~ anchor_volume"]
        value = float(auc.loc[(auc["строк"] == 160) & (auc["признак"] == feature), "AUC"].iloc[0])
        checks = [
            ("счётчик > 0 меньше чем у 40% в каждом классе",
             ", ".join(f"{c}: {v}" for c, v in zip(nonzero["класс"], nonzero["доля счётчик > 0"])),
             bool((nonzero["доля счётчик > 0"] < 0.40).all())),
            ("|Спирмен ~ anchor_volume| ≥ 0.6 в любом классе",
             ", ".join(f"{v}" for v in worst_rho), bool((worst_rho.abs() >= 0.6).any())),
            ("AUC одного признака в 0.45–0.55", f"{value}", 0.45 <= value <= 0.55)]
        if feature == "share_patent":
            checks.append(("больше 5% строк в потолке total", "потолок не обнаружен (П0.1): 0", False))
        rows += [{"признак": feature, "условие": name, "значение": shown, "выполнено": hit}
                 for name, shown, hit in checks]
        rows.append({"признак": feature, "условие": "ИТОГ: остановлен", "значение": "",
                     "выполнено": any(hit for _, _, hit in checks)})
    return pd.DataFrame(rows)


def sheet_names() -> list[str]:
    """Имена листов report_tables.xlsx; книга закрывается сразу, иначе Windows не даст дописать."""
    if not REPORT.exists():
        return []
    book = load_workbook(REPORT, read_only=True)
    names = list(book.sheetnames)
    book.close()
    return names


def append_sheets(sheets: dict[str, pd.DataFrame]) -> list[str]:
    """Дописывает листы в report_tables.xlsx; прежние листы обязаны остаться на месте."""
    before = sheet_names()
    clash = [name for name in sheets if name[:31] in before and not name.startswith("П")]
    if clash:
        raise ValueError(f"листы {clash} уже есть и не относятся к задаче П")
    mode = {"mode": "a", "if_sheet_exists": "replace"} if REPORT.exists() else {"mode": "w"}
    with pd.ExcelWriter(REPORT, engine="openpyxl", **mode) as writer:
        for title, table in sheets.items():
            table.to_excel(writer, sheet_name=title[:31], index=False)
    after = sheet_names()
    lost = [name for name in before if name not in after]
    if lost:
        raise RuntimeError(f"пропали прежние листы: {lost}")
    return after


def tables() -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    """Все таблицы П0 и таблица признаков для П1."""
    frame = feature_table()
    frames = subsets(frame)
    desc, auc, rho = describe(frames), auc_table(frames), spearman_table(frames)
    sheets = {**probe_sheets(), "П0 окна": windows(), "П0 описание": desc,
              "П0 AUC": auc, "П0 Спирмен": rho, "П0 AUC по терцилям": tertile_auc(frames),
              "П0 ru ~ patent": cross_spearman(frame), "П0 правило остановки":
              stop_rule(desc, auc, rho)}
    return sheets, frame


def main() -> None:
    """Считает П0, печатает таблицы и дописывает листы в Excel."""
    sheets, _ = tables()
    for title, table in sheets.items():
        print(f"\n{title}\n{table.to_string(index=False)}")
    names = append_sheets(sheets)
    print(f"\nлистов в {REPORT.name}: {len(names)}: {names}")


if __name__ == "__main__":
    main()
